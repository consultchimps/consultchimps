"""Write neutral, generated input files for one skills eval case into the
current directory. Usage: python make-fixtures.py <set>. Needs openpyxl."""

import csv
import random
import sys
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook

rng = random.Random(7)
PRODUCTS = ["Widget", "Gadget", "Bracket", "Fastener", "Hinge", "Gasket"]
REGIONS = ["North", "South", "East", "West"]


def save(wb, path):
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def price_rows(n, prefix):
    start = date(2026, 3, 1)
    return [
        [
            f"{prefix}-{100 + i}",
            rng.choice(PRODUCTS),
            round(rng.uniform(2, 90), 2),
            rng.randint(5, 400),
            start + timedelta(days=rng.randint(0, 60)),
        ]
        for i in range(n)
    ]


def suppliers():
    standard = ["Item Code", "Description", "Unit Price", "Quantity", "Delivery Date"]
    layouts = {
        "supplier-a": (standard, 12, []),
        "supplier-b": ([h.lower() for h in standard], 9, []),
        "supplier-c": (standard, 11, [["Supplier C price submission"], ["Prepared by: sales desk"]]),
        "supplier-d": (["SKU", "Description", "Price", "Qty", "Delivery Date"], 14, []),
        "supplier-e": (standard + ["Notes"], 8, []),
    }
    for name, (headers, n, title) in layouts.items():
        wb = Workbook()
        ws = wb.active
        ws.title = "Prices"
        for row in title:
            ws.append(row)
        ws.append(headers)
        for row in price_rows(n, name[-1].upper()):
            ws.append(row + (["check lead time"] if "Notes" in headers else []))
        save(wb, f"suppliers/{name}.xlsx")


def sales_rows():
    rows = []
    for month in range(1, 7):
        for region in REGIONS:
            for product in rng.sample(PRODUCTS, 3):
                rows.append([date(2026, month, 1), region, product, rng.randint(20, 500), round(rng.uniform(4, 60), 2)])
    return rows


def sales():
    wb = Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["Month", "Region", "Product", "Units", "Unit Price", "Revenue"])
    for i, row in enumerate(sales_rows(), start=2):
        ws.append(row + [f"=D{i}*E{i}"])
    save(wb, "sales.xlsx")


def regions():
    for region in ["North", "South", "West"]:
        wb = Workbook()
        ws = wb.active
        ws.title = "Sales"
        ws.append(["Month", "Region", "Product", "Units", "Revenue"])
        for row in sales_rows():
            if row[1] == region:
                ws.append([row[0], row[1], row[2], row[3], round(row[3] * row[4], 2)])
        save(wb, f"regions/{region.lower()}.xlsx")


def ledger():
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    summary.append(["General ledger extract"])
    summary.append([])
    summary.append(["Account", "Debit", "Credit"])
    tx = wb.create_sheet("Transactions")
    tx.append(["Date", "Account", "Reference", "Debit", "Credit"])
    accounts = ["Cash", "Receivables", "Payables", "Revenue", "Rent", "Payroll"]
    for i in range(180):
        amount = round(rng.uniform(50, 5000), 2)
        debit = rng.random() < 0.5
        tx.append([date(2026, 1, 1) + timedelta(days=i % 90), rng.choice(accounts), f"JV-{2000 + i}",
                   amount if debit else None, None if debit else amount])
    for account in accounts:
        summary.append([account, f'=SUMIFS(Transactions!D:D,Transactions!B:B,A{summary.max_row + 1})',
                        f'=SUMIFS(Transactions!E:E,Transactions!B:B,A{summary.max_row + 1})'])
    lookup = wb.create_sheet("Lookup")
    lookup.sheet_state = "hidden"
    lookup.append(["Account", "Owner"])
    for account in accounts:
        lookup.append([account, rng.choice(["Finance", "Operations"])])
    save(wb, "data/ledger.xlsx")


def pdf(path, pages):
    """A minimal valid PDF with one line of text per page."""
    objects = ["<< /Type /Catalog /Pages 2 0 R >>", None, "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    kids = []
    for text in pages:
        stream = f"BT /F1 24 Tf 72 720 Td ({text}) Tj ET"
        objects.append(f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream")
        objects.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {len(objects)} 0 R "
                       "/Resources << /Font << /F1 3 0 R >> >> >>")
        kids.append(f"{len(objects)} 0 R")
    objects[1] = f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {len(kids)} >>"
    out, offsets = "%PDF-1.4\n", []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n{body}\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n" + "".join(f"{o:010d} 00000 n \n" for o in offsets)
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n"
    Path(path).write_bytes(out.encode("latin-1"))


def pdfs():
    pdf("q1.pdf", ["Quarter 1 update", "Quarter 1 appendix"])
    pdf("q2.pdf", ["Quarter 2 update", "Quarter 2 appendix"])


def contacts():
    with open("contacts.csv", "w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(["name", "email", "team"])
        for i in range(10):
            writer.writerow([f"Person {i + 1}", f"person{i + 1}@example.com", rng.choice(["Ops", "Finance"])])


def orders():
    wb = Workbook()
    ws = wb.active
    ws.title = "Orders"
    ws.append(["Order", "Product", "Qty", "Unit Price"])
    for i in range(25):
        ws.append([f"SO-{500 + i}", rng.choice(PRODUCTS), rng.randint(1, 50), round(rng.uniform(3, 80), 2)])
    save(wb, "orders.xlsx")


SETS = {"suppliers": suppliers, "sales": sales, "regions": regions, "ledger": ledger,
        "pdfs": pdfs, "contacts": contacts, "orders": orders, "none": lambda: None}

if __name__ == "__main__":
    for name in sys.argv[1:]:
        SETS[name]()
