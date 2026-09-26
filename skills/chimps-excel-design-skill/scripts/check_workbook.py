"""Check an authored workbook against the machine-checkable part of the
delivery checklist (references/delivery-checklist.md).

Usage:
    python scripts/check_workbook.py model.xlsx [--json]

Needs only Python 3.9+ and openpyxl, which ChatGPT and Claude code sandboxes
ship with. It reads the file and changes nothing.

Each finding is FAIL (breaks the checklist), REVIEW (needs a person to judge)
or INFO (for the handover). The exit code is 1 when anything FAILs.

What it cannot check, and so leaves to the checklist: invented figures, sources
and owners on Assumptions, hardcoded figures inside calculation ranges, visible
error values (a library writes no cached results; open the file in Excel), and
the handover itself.
"""

from __future__ import annotations

import json
import re
import sys
from collections import defaultdict

try:
    import openpyxl
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover
    sys.exit("openpyxl is required: pip install openpyxl")

# Functions that must be written with a prefix, and the prefix each needs.
# Kept in step with references/formula-prefixes.md.
XLWS = {"FILTER", "SORT"}
XLFN = {
    "XLOOKUP", "XMATCH", "SORTBY", "UNIQUE", "SEQUENCE", "RANDARRAY", "LET",
    "LAMBDA", "ISOMITTED", "TEXTSPLIT", "TEXTBEFORE", "TEXTAFTER", "VSTACK",
    "HSTACK", "TOCOL", "TOROW", "CHOOSECOLS", "CHOOSEROWS", "BYROW", "BYCOL",
    "MAP", "REDUCE", "SCAN", "MAKEARRAY", "IFS", "SWITCH", "MAXIFS", "MINIFS",
    "CONCAT", "TEXTJOIN", "ANCHORARRAY", "SINGLE",
}
# Functions Excel 2007 already had: a prefix on one of these is also #NAME?.
LEGACY = {
    "SUM", "IF", "INDEX", "MATCH", "VLOOKUP", "HLOOKUP", "SUMIFS", "SUMIF",
    "COUNTIFS", "COUNTIF", "AVERAGEIFS", "IFERROR", "AND", "OR", "ROUND",
    "MAX", "MIN", "AVERAGE", "COUNT", "COUNTA", "LOOKUP", "OFFSET", "TEXT",
}
LIBRARY_NAMES = ("openpyxl", "sheetjs", "xlsx", "xlsxwriter", "python",
                 "pandas", "exceljs", "apache poi", "epplus")
BLUE, BLACK, GREEN = "0000FF", "000000", "008000"
DATE_TEXT = re.compile(
    r"^\s*(\d{1,2}[/.-]\d{1,2}[/.-]\d{2,4}|\d{4}-\d{2}-\d{2})\s*$")
CELL_REF = re.compile(r"^\$?[A-Z]{1,3}\$?\d+$")

findings: list[dict[str, str]] = []


def add(level: str, rule: str, where: str, detail: str) -> None:
    findings.append(
        {"level": level, "rule": rule, "where": where, "detail": detail})


def rgb(font) -> str | None:
    color = font.color if font else None
    if color is None or color.type != "rgb" or not isinstance(color.rgb, str):
        return None
    return color.rgb[-6:].upper()


def strip_strings(formula: str) -> str:
    return re.sub(r'"[^"]*"', '""', formula)


def check_formula(where: str, formula: str, lookups: dict[str, list[str]],
                  names_used: set[str]) -> None:
    body = strip_strings(formula)
    for match in re.finditer(r"((?:_xlfn\.)?(?:_xlws\.)?)([A-Z][A-Z0-9.]*)\(",
                             body):
        prefix, name = match.group(1), match.group(2)
        if name in XLWS and prefix != "_xlfn._xlws.":
            add("FAIL", "prefix", where,
                f"{name} must be written _xlfn._xlws.{name}")
        elif name in XLFN and prefix != "_xlfn.":
            add("FAIL", "prefix", where, f"{name} must be written _xlfn.{name}")
        elif name in LEGACY and prefix:
            add("FAIL", "prefix", where,
                f"{name} existed in Excel 2007 and takes no prefix")
        if name in {"XLOOKUP", "XMATCH"}:
            lookups["XLOOKUP"].append(where)
        elif name in {"VLOOKUP", "HLOOKUP", "INDEX", "MATCH"}:
            lookups["INDEX/MATCH"].append(where)
    if re.search(r"#(REF|NAME|VALUE|DIV/0|N/A)[!?]?", body):
        add("FAIL", "error-in-formula", where, "the formula text holds an error")
    if re.search(r"(?:^|[^A-Za-z_])'?Assumptions'?!\$?[A-Z]{1,3}\$?\d+", body):
        add("FAIL", "named-range", where,
            "points at an Assumptions cell by address; use its named range")
    for token in re.findall(r"(?<![A-Za-z0-9_.!$'])([A-Za-z_][A-Za-z0-9_.]*)"
                            r"(?![A-Za-z0-9_.(!])", body):
        upper = token.upper()
        if (CELL_REF.match(upper) or upper in {"TRUE", "FALSE"}
                or token.startswith("_xl")):
            continue
        if "_" in token:
            names_used.add(token)


def main(path: str, as_json: bool) -> int:
    wb = openpyxl.load_workbook(path)
    props = wb.properties
    for field in ("creator", "lastModifiedBy"):
        value = (getattr(props, field) or "").strip()
        if not value:
            add("FAIL", "properties", "workbook", f"{field} is not set")
        elif any(lib in value.lower() for lib in LIBRARY_NAMES):
            add("FAIL", "properties", "workbook",
                f"{field} reads '{value}', a library name")
    if not (props.title or "").strip():
        add("FAIL", "properties", "workbook", "title is not set")

    defined = {name.upper() for name in wb.defined_names}
    for ws in wb.worksheets:
        defined.update(name.upper() for name in ws.defined_names)
    names_used: set[str] = set()
    lookups: dict[str, list[str]] = defaultdict(list)

    for ws in wb.worksheets:
        sheet = ws.title
        if ws.max_row < 2:
            continue
        is_data = ws.max_row > 2 and sheet.lower() != "assumptions"

        for rng in ws.merged_cells.ranges:
            if rng.min_row > 1:
                add("FAIL", "merged-cells", f"{sheet}!{rng.coord}",
                    "merged outside the title row")
        if is_data and not ws.freeze_panes:
            add("FAIL", "freeze", sheet, "header row is not frozen")
        if is_data and not ws.auto_filter.ref:
            add("FAIL", "autofilter", sheet, "autofilter is not on")

        formats: dict[int, set[str]] = defaultdict(set)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                where = f"{sheet}!{cell.coordinate}"
                value = cell.value
                if cell.comment is not None:
                    add("REVIEW", "comments", where,
                        "cell comment; keep only if the user asked for it")
                if value is None:
                    continue
                colour = rgb(cell.font)
                if isinstance(value, str) and value.startswith("="):
                    check_formula(where, value, lookups, names_used)
                    want = GREEN if "!" in strip_strings(value) else BLACK
                    if colour not in (None, want) or (
                            colour is None and want == GREEN):
                        add("FAIL", "font-colour", where,
                            f"formula font should be {want}, is {colour}")
                elif isinstance(value, (int, float)) and not isinstance(
                        value, bool):
                    if colour != BLUE:
                        add("REVIEW", "font-colour", where,
                            f"hardcoded number in {colour or 'default'}; an "
                            f"input is blue {BLUE}, a figure in a calculation "
                            "range should be a formula")
                elif isinstance(value, str):
                    if DATE_TEXT.match(value):
                        add("FAIL", "date-as-text", where,
                            f"'{value}' is text that looks like a date")
                    if value.strip().lower() == "needs input":
                        add("INFO", "needs-input", where,
                            "list this in the handover")
                # Text labels carry no number format worth comparing.
                if not isinstance(value, str) or value.startswith("="):
                    formats[cell.column].add(cell.number_format)
        for column, found in formats.items():
            if len(found) > 1:
                add("FAIL", "number-format",
                    f"{sheet}!{get_column_letter(column)}",
                    "mixes formats: " + ", ".join(sorted(found)))

    if lookups["XLOOKUP"] and lookups["INDEX/MATCH"]:
        add("FAIL", "lookup-family", "workbook",
            f"XLOOKUP ({lookups['XLOOKUP'][0]}) and INDEX/MATCH or VLOOKUP "
            f"({lookups['INDEX/MATCH'][0]}) are both used")
    for name in sorted(names_used):
        if name.upper() not in defined:
            add("REVIEW", "defined-names", "workbook",
                f"'{name}' reads like a named range but is not defined; "
                "an unregistered name shows #NAME?")

    failed = sum(f["level"] == "FAIL" for f in findings)
    if as_json:
        print(json.dumps({"file": path, "failed": failed,
                          "findings": findings}, indent=2))
    else:
        for level in ("FAIL", "REVIEW", "INFO"):
            for f in (f for f in findings if f["level"] == level):
                print(f"{level:6} {f['rule']:14} {f['where']:22} {f['detail']}")
        print(f"\n{failed} failing, "
              f"{sum(f['level'] == 'REVIEW' for f in findings)} to review. "
              "The rest of references/delivery-checklist.md is still yours "
              "to apply.")
    return 1 if failed else 0


if __name__ == "__main__":
    args = [a for a in sys.argv[1:] if a != "--json"]
    if len(args) != 1:
        sys.exit(__doc__)
    sys.exit(main(args[0], "--json" in sys.argv[1:]))
